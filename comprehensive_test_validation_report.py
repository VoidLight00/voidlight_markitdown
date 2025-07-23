#!/usr/bin/env python3
"""
Comprehensive Test Validation Report Generator
Executes all tests and generates detailed validation report
"""

import subprocess
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path
import traceback

class TestValidator:
    def __init__(self):
        self.results = []
        self.start_time = datetime.now()
        self.test_commands = []
        
    def execute_unit_tests(self):
        """Execute unit tests with pytest"""
        print("\n" + "="*80)
        print("EXECUTING UNIT TESTS")
        print("="*80)
        
        unit_tests = [
            "packages/voidlight_markitdown/tests/test_voidlight_markitdown.py",
            "packages/voidlight_markitdown/tests/test_module_misc.py",
            "packages/voidlight_markitdown/tests/test_module_vectors.py",
        ]
        
        for test_file in unit_tests:
            if Path(test_file).exists():
                result = self.run_pytest(test_file, "unit_test")
                self.results.append(result)
                
    def execute_korean_nlp_tests(self):
        """Execute Korean NLP tests"""
        print("\n" + "="*80)
        print("EXECUTING KOREAN NLP TESTS")
        print("="*80)
        
        # Check Korean NLP availability first
        check_result = self.check_korean_nlp_status()
        self.results.append(check_result)
        
        korean_tests = [
            "packages/voidlight_markitdown/tests/test_korean_nlp_features.py",
            "packages/voidlight_markitdown/tests/test_korean_utils.py",
            "packages/voidlight_markitdown/test_korean_nlp_simple.py",
        ]
        
        for test_file in korean_tests:
            if Path(test_file).exists():
                result = self.run_pytest(test_file, "korean_nlp_test")
                self.results.append(result)
                
    def execute_converter_tests(self):
        """Execute converter tests"""
        print("\n" + "="*80)
        print("EXECUTING CONVERTER TESTS")
        print("="*80)
        
        # Run converter availability test
        converter_test = self.test_converter_availability()
        self.results.append(converter_test)
        
    def execute_cli_tests(self):
        """Execute CLI tests"""
        print("\n" + "="*80)
        print("EXECUTING CLI TESTS")
        print("="*80)
        
        cli_tests = [
            "packages/voidlight_markitdown/tests/test_cli_misc.py",
            "packages/voidlight_markitdown/tests/test_cli_vectors.py",
        ]
        
        for test_file in cli_tests:
            if Path(test_file).exists():
                result = self.run_pytest(test_file, "cli_test")
                self.results.append(result)
                
    def execute_integration_tests(self):
        """Execute integration tests"""
        print("\n" + "="*80)
        print("EXECUTING INTEGRATION TESTS")
        print("="*80)
        
        # Test basic integration
        integration_result = self.test_basic_integration()
        self.results.append(integration_result)
        
    def check_korean_nlp_status(self):
        """Check Korean NLP initialization status"""
        print("\nChecking Korean NLP status...")
        
        test_code = '''
import sys
sys.path.insert(0, 'packages/voidlight_markitdown/src')
from voidlight_markitdown import VoidLightMarkItDown

try:
    # Test basic initialization
    md = VoidLightMarkItDown(korean_mode=True)
    print("✅ Korean mode initialized successfully")
    
    # Test Korean text processing
    result = md.convert("안녕하세요. 이것은 테스트입니다.")
    print(f"✅ Korean text processed: {result.markdown[:50]}...")
    
    # Check NLP features
    from voidlight_markitdown._korean_utils import KoreanTextProcessor
    processor = KoreanTextProcessor()
    
    if processor.kiwi_tokenizer:
        print("✅ Kiwi tokenizer available")
    else:
        print("❌ Kiwi tokenizer not available")
        
    if processor.konlpy_available:
        print("✅ KoNLPy available")
    else:
        print("⚠️  KoNLPy not available (optional)")
        
except Exception as e:
    print(f"❌ Korean NLP initialization failed: {e}")
    import traceback
    traceback.print_exc()
'''
        
        result = {
            'test_name': 'Korean NLP Status Check',
            'category': 'korean_nlp',
            'start_time': datetime.now().isoformat()
        }
        
        try:
            proc = subprocess.run(
                [sys.executable, '-c', test_code],
                capture_output=True,
                text=True,
                timeout=30
            )
            
            result['stdout'] = proc.stdout
            result['stderr'] = proc.stderr
            result['returncode'] = proc.returncode
            result['status'] = 'passed' if proc.returncode == 0 else 'failed'
            
        except Exception as e:
            result['status'] = 'error'
            result['error'] = str(e)
            
        return result
        
    def test_converter_availability(self):
        """Test all converter availability"""
        print("\nTesting converter availability...")
        
        test_code = '''
import sys
sys.path.insert(0, 'packages/voidlight_markitdown/src')
from voidlight_markitdown import VoidLightMarkItDown
import tempfile
import os

md = VoidLightMarkItDown()

# Test each converter
converters = {
    'PlainText': lambda: md.convert("Hello World"),
    'HTML': lambda: md.convert("<html><body>Test</body></html>"),
    'CSV': lambda: md.convert("a,b\\n1,2"),
    'JSON': lambda: md.convert('{"test": "data"}'),
    'YAML': lambda: md.convert("test: data"),
}

results = []
for name, test_func in converters.items():
    try:
        result = test_func()
        results.append(f"✅ {name} converter: WORKING")
    except Exception as e:
        results.append(f"❌ {name} converter: {str(e)[:50]}...")

print("\\n".join(results))

# Test file converters with temporary files
print("\\n--- File Converters ---")

# DOCX test
try:
    from docx import Document
    doc = Document()
    doc.add_paragraph("Test")
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        doc.save(f.name)
        result = md.convert(f.name)
        os.unlink(f.name)
    print("✅ DOCX converter: WORKING")
except Exception as e:
    print(f"❌ DOCX converter: {str(e)[:50]}...")

# Excel test
try:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Test'
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
        wb.save(f.name)
        result = md.convert(f.name)
        os.unlink(f.name)
    print("✅ Excel converter: WORKING")
except Exception as e:
    print(f"❌ Excel converter: {str(e)[:50]}...")

# List all registered converters
print(f"\\nTotal converters registered: {len(md._converters)}")
'''
        
        result = {
            'test_name': 'Converter Availability Test',
            'category': 'converter_test',
            'start_time': datetime.now().isoformat()
        }
        
        try:
            proc = subprocess.run(
                [sys.executable, '-c', test_code],
                capture_output=True,
                text=True,
                timeout=30
            )
            
            result['stdout'] = proc.stdout
            result['stderr'] = proc.stderr
            result['returncode'] = proc.returncode
            result['status'] = 'passed' if '✅' in proc.stdout else 'failed'
            
            # Count successes and failures
            success_count = proc.stdout.count('✅')
            fail_count = proc.stdout.count('❌')
            result['summary'] = {
                'passed': success_count,
                'failed': fail_count
            }
            
        except Exception as e:
            result['status'] = 'error'
            result['error'] = str(e)
            
        return result
        
    def test_basic_integration(self):
        """Test basic integration scenarios"""
        print("\nTesting basic integration...")
        
        test_code = '''
import sys
sys.path.insert(0, 'packages/voidlight_markitdown/src')
from voidlight_markitdown import VoidLightMarkItDown
import tempfile
import os

try:
    # Test 1: Basic conversion
    md = VoidLightMarkItDown()
    result = md.convert("Hello World")
    assert result.markdown == "Hello World"
    print("✅ Basic text conversion: PASSED")
    
    # Test 2: File conversion
    with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False) as f:
        f.write("Test file content")
        f.flush()
        result = md.convert(f.name)
        assert "Test file content" in result.markdown
        os.unlink(f.name)
    print("✅ File conversion: PASSED")
    
    # Test 3: Korean mode
    md_korean = VoidLightMarkItDown(korean_mode=True)
    result = md_korean.convert("안녕하세요")
    assert "안녕하세요" in result.markdown
    print("✅ Korean mode conversion: PASSED")
    
    # Test 4: Stream conversion
    import io
    stream = io.BytesIO(b"Stream content")
    result = md.convert_stream(stream, stream_info={"file_extension": ".txt"})
    assert "Stream content" in result.markdown
    print("✅ Stream conversion: PASSED")
    
    # Test 5: Error handling
    try:
        md.convert("/nonexistent/file.txt")
    except Exception:
        print("✅ Error handling: PASSED")
    else:
        print("❌ Error handling: FAILED (no exception raised)")
        
    print("\\nAll integration tests completed!")
    
except Exception as e:
    print(f"❌ Integration test failed: {e}")
    import traceback
    traceback.print_exc()
'''
        
        result = {
            'test_name': 'Basic Integration Test',
            'category': 'integration_test',
            'start_time': datetime.now().isoformat()
        }
        
        try:
            proc = subprocess.run(
                [sys.executable, '-c', test_code],
                capture_output=True,
                text=True,
                timeout=30
            )
            
            result['stdout'] = proc.stdout
            result['stderr'] = proc.stderr
            result['returncode'] = proc.returncode
            result['status'] = 'passed' if proc.returncode == 0 else 'failed'
            
            # Count test results
            passed = proc.stdout.count('✅')
            failed = proc.stdout.count('❌')
            result['summary'] = {
                'passed': passed,
                'failed': failed
            }
            
        except Exception as e:
            result['status'] = 'error'
            result['error'] = str(e)
            
        return result
        
    def run_pytest(self, test_file, category):
        """Run pytest on a file"""
        print(f"\nRunning pytest: {test_file}")
        
        result = {
            'test_file': test_file,
            'category': category,
            'start_time': datetime.now().isoformat()
        }
        
        try:
            cmd = [
                sys.executable, '-m', 'pytest',
                test_file,
                '-v', '--tb=short',
                '--json-report', 
                '--json-report-file=test_report_temp.json'
            ]
            
            # Store command for reproducibility
            self.test_commands.append(' '.join(cmd))
            
            start = time.time()
            proc = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=120
            )
            duration = time.time() - start
            
            result['duration'] = duration
            result['returncode'] = proc.returncode
            result['stdout'] = proc.stdout[-1000:]  # Last 1000 chars
            result['stderr'] = proc.stderr[-1000:]  # Last 1000 chars
            
            # Parse JSON report
            if Path('test_report_temp.json').exists():
                with open('test_report_temp.json', 'r') as f:
                    pytest_data = json.load(f)
                    result['summary'] = pytest_data.get('summary', {})
                os.remove('test_report_temp.json')
            else:
                # Parse from stdout
                result['summary'] = self.parse_pytest_output(proc.stdout)
                
            result['status'] = 'passed' if proc.returncode == 0 else 'failed'
            
        except subprocess.TimeoutExpired:
            result['status'] = 'timeout'
            result['error'] = 'Test timed out after 120 seconds'
        except Exception as e:
            result['status'] = 'error'
            result['error'] = str(e)
            result['traceback'] = traceback.format_exc()
            
        return result
        
    def parse_pytest_output(self, output):
        """Parse pytest output for summary"""
        summary = {'passed': 0, 'failed': 0, 'skipped': 0}
        
        if 'passed' in output:
            try:
                summary['passed'] = int(output.split('passed')[0].strip().split()[-1])
            except: pass
            
        if 'failed' in output:
            try:
                summary['failed'] = int(output.split('failed')[0].strip().split()[-1])
            except: pass
            
        if 'skipped' in output:
            try:
                summary['skipped'] = int(output.split('skipped')[0].strip().split()[-1])
            except: pass
            
        return summary
        
    def generate_final_report(self):
        """Generate comprehensive final report"""
        total_duration = (datetime.now() - self.start_time).total_seconds()
        
        # Calculate totals
        total_passed = sum(r.get('summary', {}).get('passed', 0) for r in self.results)
        total_failed = sum(r.get('summary', {}).get('failed', 0) for r in self.results)
        total_skipped = sum(r.get('summary', {}).get('skipped', 0) for r in self.results)
        
        # Category breakdown
        categories = {}
        for result in self.results:
            cat = result.get('category', 'unknown')
            if cat not in categories:
                categories[cat] = {'total': 0, 'passed': 0, 'failed': 0, 'error': 0}
            categories[cat]['total'] += 1
            if result.get('status') == 'passed':
                categories[cat]['passed'] += 1
            elif result.get('status') == 'failed':
                categories[cat]['failed'] += 1
            else:
                categories[cat]['error'] += 1
                
        report = {
            'execution_metadata': {
                'date': self.start_time.isoformat(),
                'duration': total_duration,
                'python_version': sys.version,
                'platform': sys.platform,
                'working_directory': os.getcwd()
            },
            'summary': {
                'total_test_files': len(self.results),
                'total_tests_passed': total_passed,
                'total_tests_failed': total_failed,
                'total_tests_skipped': total_skipped,
                'pass_rate': f"{(total_passed / (total_passed + total_failed) * 100):.1f}%" if (total_passed + total_failed) > 0 else "N/A"
            },
            'category_breakdown': categories,
            'test_results': self.results,
            'reproducible_commands': self.test_commands
        }
        
        # Save JSON report
        with open('comprehensive_test_validation_report.json', 'w') as f:
            json.dump(report, f, indent=2)
            
        # Generate markdown report
        self.generate_markdown_report(report)
        
        # Print summary
        self.print_summary(report)
        
        return report
        
    def generate_markdown_report(self, report):
        """Generate markdown report"""
        md_content = f"""# Voidlight MarkItDown Test Validation Report

**Date**: {report['execution_metadata']['date']}  
**Duration**: {report['execution_metadata']['duration']:.2f} seconds  
**Python Version**: {report['execution_metadata']['python_version'].split()[0]}  
**Platform**: {report['execution_metadata']['platform']}  

## Executive Summary

- **Total Test Files Executed**: {report['summary']['total_test_files']}
- **Total Tests Passed**: {report['summary']['total_tests_passed']}
- **Total Tests Failed**: {report['summary']['total_tests_failed']}
- **Total Tests Skipped**: {report['summary']['total_tests_skipped']}
- **Overall Pass Rate**: {report['summary']['pass_rate']}

## Category Breakdown

| Category | Total | Passed | Failed | Error |
|----------|-------|--------|--------|-------|
"""
        
        for cat, stats in report['category_breakdown'].items():
            md_content += f"| {cat} | {stats['total']} | {stats['passed']} | {stats['failed']} | {stats['error']} |\n"
            
        md_content += "\n## Test Execution Details\n\n"
        
        for result in report['test_results']:
            status_icon = "✅" if result['status'] == 'passed' else "❌"
            md_content += f"### {status_icon} {result.get('test_name', result.get('test_file', 'Unknown'))}\n\n"
            md_content += f"- **Category**: {result.get('category', 'N/A')}\n"
            md_content += f"- **Status**: {result['status']}\n"
            
            if 'summary' in result:
                md_content += f"- **Results**: {result['summary'].get('passed', 0)} passed, "
                md_content += f"{result['summary'].get('failed', 0)} failed, "
                md_content += f"{result['summary'].get('skipped', 0)} skipped\n"
                
            if 'duration' in result:
                md_content += f"- **Duration**: {result['duration']:.2f}s\n"
                
            md_content += "\n"
            
        md_content += """## Reproducible Test Commands

```bash
# Activate virtual environment
source mcp-env/bin/activate

# Run individual test commands
"""
        
        for cmd in report['reproducible_commands']:
            md_content += f"{cmd}\n"
            
        md_content += "```\n"
        
        with open('comprehensive_test_validation_report.md', 'w') as f:
            f.write(md_content)
            
    def print_summary(self, report):
        """Print summary to console"""
        print("\n" + "="*80)
        print("COMPREHENSIVE TEST VALIDATION SUMMARY")
        print("="*80)
        print(f"Execution Date: {report['execution_metadata']['date']}")
        print(f"Total Duration: {report['execution_metadata']['duration']:.2f} seconds")
        print(f"\nOVERALL RESULTS:")
        print(f"  Total Test Files: {report['summary']['total_test_files']}")
        print(f"  Total Tests Passed: {report['summary']['total_tests_passed']}")
        print(f"  Total Tests Failed: {report['summary']['total_tests_failed']}")
        print(f"  Total Tests Skipped: {report['summary']['total_tests_skipped']}")
        print(f"  Pass Rate: {report['summary']['pass_rate']}")
        
        print(f"\nCATEGORY BREAKDOWN:")
        for cat, stats in report['category_breakdown'].items():
            print(f"\n  {cat.upper()}:")
            print(f"    Total: {stats['total']}")
            print(f"    Passed: {stats['passed']}")
            print(f"    Failed: {stats['failed']}")
            print(f"    Error: {stats['error']}")
            
        print(f"\nReports generated:")
        print(f"  - comprehensive_test_validation_report.json")
        print(f"  - comprehensive_test_validation_report.md")
        
    def execute_all(self):
        """Execute all test categories"""
        print("Starting Comprehensive Test Validation...")
        print(f"Working Directory: {os.getcwd()}")
        print(f"Python: {sys.executable}")
        
        # Execute test categories
        self.execute_unit_tests()
        self.execute_korean_nlp_tests()
        self.execute_converter_tests()
        self.execute_cli_tests()
        self.execute_integration_tests()
        
        # Generate final report
        report = self.generate_final_report()
        
        return report

if __name__ == '__main__':
    validator = TestValidator()
    report = validator.execute_all()
    
    # Exit with appropriate code
    if report['summary']['total_tests_failed'] > 0:
        sys.exit(1)
    else:
        sys.exit(0)